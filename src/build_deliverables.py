#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Builds presentation deliverables from the model outputs:
  * outputs/figures/*.png   (publication-quality charts)
  * outputs/WorldCup2026_Forecast.xlsx   (multi-sheet professional workbook)

Run AFTER worldcup2026_model.py.  Reads only the CSVs in outputs/.
"""
import os, datetime
import numpy as np, pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT  = os.path.join(ROOT, "outputs"); FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)

# palette
NAVY="#0A1F44"; GOLD="#C9A227"; RED="#C1272D"; TEAL="#2A9D8F"; GREY="#6c757d"
plt.rcParams.update({"font.size":11,"axes.titleweight":"bold","axes.edgecolor":"#888",
                     "axes.grid":True,"grid.alpha":.25,"figure.dpi":130,
                     "axes.titlecolor":NAVY,"font.family":"DejaVu Sans"})

full = pd.read_csv(os.path.join(OUT,"forecast_full.csv"))
champ= pd.read_csv(os.path.join(OUT,"champion_probabilities.csv"))
grp  = pd.read_csv(os.path.join(OUT,"group_stage_probabilities.csv"))
appA = pd.read_csv(os.path.join(OUT,"technical_appendix_A_poisson_binomial.csv"))
mm   = pd.read_csv(os.path.join(OUT,"match_matrix_group.csv"))
cal  = pd.read_csv(os.path.join(OUT,"calibration.csv"))
sens = pd.read_csv(os.path.join(OUT,"sensitivity_form_sd.csv"))

# =========================================================================== #
#  FIGURES
# =========================================================================== #
def fig_champion():
    d = full.sort_values("P_consensus",ascending=False).head(16).iloc[::-1]
    y = np.arange(len(d)); h=0.26
    fig,ax = plt.subplots(figsize=(9.5,7.2))
    ax.barh(y+h, d["P_champion"]*100, h, label="Model (Elo+MC)", color=NAVY)
    ax.barh(y,    d["P_consensus"]*100, h, label="Consensus (blend)", color=GOLD)
    ax.barh(y-h, d["P_market"]*100, h, label="Market (de-vigged)", color=TEAL)
    ax.set_yticks(y); ax.set_yticklabels(d["team"])
    ax.set_xlabel("Probability of winning the 2026 World Cup (%)")
    ax.set_title("2026 World Cup — Title Probability: Model vs Market vs Consensus")
    ax.legend(loc="lower right", framealpha=.95)
    for i,v in zip(y, d["P_consensus"]*100):
        ax.text(v+0.2, i, f"{v:.1f}%", va="center", fontsize=8.5, color=NAVY)
    plt.tight_layout(); plt.savefig(f"{FIG}/fig1_champion_top16.png"); plt.close()

def fig_value():
    d = full[full["P_market"]>0.004].copy()
    fig,ax=plt.subplots(figsize=(8.8,7.2))
    mx = max(d["P_market"].max(), d["P_champion"].max())*100*1.08
    ax.plot([0,mx],[0,mx],"--",color=GREY,lw=1,label="model = market")
    ax.scatter(d["P_market"]*100, d["P_champion"]*100, s=70, color=NAVY, zorder=3, edgecolor="w")
    for _,r in d.iterrows():
        dv = (r["P_champion"]-r["P_market"])*100
        if abs(dv)>1.0 or r["P_consensus"]>0.05:
            ax.annotate(r["team"], (r["P_market"]*100, r["P_champion"]*100),
                        xytext=(4,4), textcoords="offset points", fontsize=8.6,
                        color=RED if dv>0 else "#1f6f8b", fontweight="bold")
    ax.set_xlabel("Market-implied title probability (%)")
    ax.set_ylabel("Model title probability (%)")
    ax.set_title("Where the model disagrees with the market\n(above line = model 'value'; below = model fade)")
    ax.legend(loc="upper left")
    plt.tight_layout(); plt.savefig(f"{FIG}/fig2_value_vs_market.png"); plt.close()

def fig_round_heatmap():
    cols=["P_reach_R32","P_reach_R16","P_reach_QF","P_reach_SF","P_reach_Final","P_champion"]
    labels=["R32","R16","QF","SF","Final","Champion"]
    d=full.sort_values("P_champion",ascending=False).head(20)
    M=d[cols].to_numpy()*100
    cmap=LinearSegmentedColormap.from_list("wc",["#ffffff","#9ec5fe",NAVY])
    fig,ax=plt.subplots(figsize=(8.5,9))
    im=ax.imshow(M,aspect="auto",cmap=cmap,vmin=0,vmax=100)
    ax.set_xticks(range(len(labels))); ax.set_xticklabels(labels)
    ax.set_yticks(range(len(d))); ax.set_yticklabels(d["team"])
    for i in range(M.shape[0]):
        for j in range(M.shape[1]):
            v=M[i,j]
            ax.text(j,i,f"{v:.0f}" if v>=1 else f"{v:.1f}",ha="center",va="center",
                    color="white" if v>55 else NAVY,fontsize=8.3)
    ax.set_title("Probability of reaching each stage (%) — top 20 teams")
    plt.colorbar(im,fraction=0.035,pad=0.02,label="%")
    plt.tight_layout(); plt.savefig(f"{FIG}/fig3_round_heatmap.png"); plt.close()

def fig_calibration():
    fig,ax=plt.subplots(figsize=(8.2,5.6))
    ax.plot(cal["elo_diff"],cal["E_elo"],color=NAVY,lw=2.2,label="Elo logistic  E = 1/(1+10^(-d/400))")
    ax.scatter(cal["elo_diff"],cal["E_model"],color=GOLD,s=28,zorder=3,label="Poisson goal model (Skellam)")
    ax.set_xlabel("Elo rating difference  d = R_A − R_B")
    ax.set_ylabel("Expected match score for A")
    ax2=ax.twinx()
    ax2.bar(cal["elo_diff"],cal["abs_err"]*1000,width=8,color=RED,alpha=.35)
    ax2.set_ylabel("|error| × 1000",color=RED); ax2.set_ylim(0,12); ax2.grid(False)
    ax.set_title("Goal model is calibrated to Elo (max error %.4f)"%cal["abs_err"].max())
    ax.legend(loc="lower right")
    plt.tight_layout(); plt.savefig(f"{FIG}/fig4_calibration.png"); plt.close()

def fig_confederation():
    c=full.groupby("confederation")["P_champion"].sum().sort_values()*100
    cm=full.groupby("confederation")["P_market"].sum().reindex(c.index)*100
    y=np.arange(len(c))
    fig,ax=plt.subplots(figsize=(8,5))
    ax.barh(y+0.2,c.values,0.4,color=NAVY,label="Model")
    ax.barh(y-0.2,cm.values,0.4,color=TEAL,label="Market")
    ax.set_yticks(y); ax.set_yticklabels(c.index)
    ax.set_xlabel("Total title probability (%)")
    ax.set_title("Where will the 2026 champion come from? (by confederation)")
    for i,v in zip(y,c.values): ax.text(v+0.4,i+0.2,f"{v:.0f}%",va="center",fontsize=9)
    ax.legend()
    plt.tight_layout(); plt.savefig(f"{FIG}/fig5_confederation.png"); plt.close()

def fig_sensitivity():
    fig,ax=plt.subplots(figsize=(8.2,5.6))
    colors=[GOLD,RED,NAVY,TEAL,"#8338ec","#fb5607"]
    for t,c in zip(["Spain","Argentina","France","England","Portugal","Brazil"],colors):
        ax.plot(sens["form_sd"],sens[t],marker="o",color=c,lw=2,label=t)
    ax.set_xlabel("Form / rating uncertainty  σ_form  (Elo)")
    ax.set_ylabel("Title probability (%)")
    ax.set_title("Sensitivity of title odds to the form-uncertainty assumption")
    ax.axvline(70,color=GREY,ls="--",lw=1); ax.text(71,ax.get_ylim()[1]*0.92,"σ=70 (used)",fontsize=8.5,color=GREY)
    ax.legend(ncol=3,fontsize=9)
    plt.tight_layout(); plt.savefig(f"{FIG}/fig6_sensitivity.png"); plt.close()

for f in (fig_champion,fig_value,fig_round_heatmap,fig_calibration,fig_confederation,fig_sensitivity):
    f(); print("figure:",f.__name__)

# =========================================================================== #
#  EXCEL WORKBOOK
# =========================================================================== #
THIN=Side(style="thin",color="D0D0D0")
BORDER=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
HEADER_FILL=PatternFill("solid",fgColor=NAVY.replace("#",""))
TITLE_FILL =PatternFill("solid",fgColor="11305f")
SUB_FILL   =PatternFill("solid",fgColor="EAEFF7")
WHITE=Font(color="FFFFFF",bold=True)
def style_header(ws,row,ncol,fill=HEADER_FILL):
    for c in range(1,ncol+1):
        cell=ws.cell(row=row,column=c); cell.fill=fill; cell.font=WHITE
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        cell.border=BORDER
def widths(ws,wmap):
    for col,w in wmap.items(): ws.column_dimensions[col].width=w
def pct(ws,rng,nf="0.00%"):
    for row in ws[rng]:
        for c in row: c.number_format=nf

wb=openpyxl.Workbook()

# ---- Sheet 1: OVERVIEW ---------------------------------------------------- #
ws=wb.active; ws.title="Overview"; ws.sheet_view.showGridLines=False
ws.merge_cells("B2:K2"); t=ws["B2"]
t.value="2026 FIFA WORLD CUP — DATA-DRIVEN CHAMPION FORECAST"
t.font=Font(size=18,bold=True,color="FFFFFF"); t.fill=TITLE_FILL
t.alignment=Alignment(horizontal="center",vertical="center"); ws.row_dimensions[2].height=30
ws.merge_cells("B3:K3"); s=ws["B3"]
s.value=("Elo → calibrated Poisson goal model → exact Poisson-binomial group stage → "
         f"{200000:,}-replication Monte-Carlo knockout.  Generated "+datetime.date.today().isoformat())
s.font=Font(size=10,italic=True,color="333333"); s.alignment=Alignment(horizontal="center")

note=[
 ("METHOD","Ratings: World Football Elo (eloratings.net, mid-June 2026). Match model calibrated so the "
          "Poisson/Skellam win-probability reproduces the Elo logistic (max error 0.4%). Group stage solved "
          "EXACTLY (3^6 enumeration + Poisson-binomial). Knockout via 200k Monte-Carlo over the official "
          "FIFA bracket incl. the best-third-place combination assignment. Rating/form uncertainty σ=70 Elo."),
 ("HEADLINE","Consensus title pick: SPAIN, narrowly over FRANCE and ARGENTINA. The pure model rates "
            "ARGENTINA and SPAIN as co-favourites; the market prefers FRANCE (opening-match form)."),
 ("MODEL vs MARKET","Biggest model 'value': ARGENTINA (model ~22% vs market ~9%) and COLOMBIA. Biggest "
                   "model 'fade': PORTUGAL (model ~5% vs market ~9%). Both directions match Nate Silver's "
                   "PELE model (bullish Argentina, bearish Portugal) — independent corroboration."),
 ("HOW TO READ","Each sheet is self-contained. 'Champion Forecast' = full 48-team ranking. "
               "'Appendix A' = the exact per-team Poisson-binomial demonstration. Colour scales: green = higher."),
]
r=5
for h,txt in note:
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=11)
    c=ws.cell(row=r,column=2,value=f"{h}:  {txt}")
    c.alignment=Alignment(wrap_text=True,vertical="top"); c.font=Font(size=10)
    c.fill=SUB_FILL if (r//2)%2 else PatternFill("solid",fgColor="FFFFFF")
    ws.row_dimensions[r].height=46; r+=1

# headline mini-table
r+=1; ws.cell(row=r,column=2,value="TOP 10 — TITLE PROBABILITY").font=Font(bold=True,size=12,color="0A1F44")
r+=1; hdr=["Rank","Team","Group","Model %","Market %","Consensus %"]
for j,hh in enumerate(hdr): ws.cell(row=r,column=2+j,value=hh)
style_header(ws,r,7); hr=r
top=full.sort_values("P_consensus",ascending=False).head(10)
for _,row_ in top.iterrows():
    r+=1
    vals=[int(row_["rank"]),row_["team"],row_["group"],row_["P_champion"],row_["P_market"],row_["P_consensus"]]
    for j,v in enumerate(vals):
        cell=ws.cell(row=r,column=2+j,value=v); cell.border=BORDER
        if j>=3: cell.number_format="0.0%"
        cell.alignment=Alignment(horizontal="center" if j!=1 else "left")
ws.conditional_formatting.add(f"E{hr+1}:G{r}",
    ColorScaleRule(start_type="min",start_color="FFFFFF",end_type="max",end_color="63BE7B"))
widths(ws,{"A":2,"B":10,"C":18,"D":9,"E":11,"F":11,"G":13,"H":3,"I":9,"J":9,"K":9})
img=XLImage(f"{FIG}/fig1_champion_top16.png"); img.width=560; img.height=420
ws.add_image(img,"M2")

# ---- generic dataframe sheet helper -------------------------------------- #
def add_df_sheet(name,df,pct_cols=(),int_cols=(),scale_cols=(),width_map=None,
                 title=None,freeze="A2",img=None,img_anchor=None):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    start=1
    if title:
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(df.columns))
        c=ws.cell(row=1,column=1,value=title); c.font=Font(size=13,bold=True,color="FFFFFF")
        c.fill=TITLE_FILL; c.alignment=Alignment(horizontal="center"); ws.row_dimensions[1].height=22
        start=2
    for j,col in enumerate(df.columns,1):
        ws.cell(row=start,column=j,value=col)
    style_header(ws,start,len(df.columns))
    for i,(_,row_) in enumerate(df.iterrows()):
        for j,col in enumerate(df.columns,1):
            v=row_[col]
            if pd.isna(v): v=None
            cell=ws.cell(row=start+1+i,column=j,value=v); cell.border=BORDER
            cell.alignment=Alignment(horizontal="left" if j<=2 else "center")
            if col in pct_cols: cell.number_format="0.0%"
            if col in int_cols and v is not None: cell.number_format="0"
    last=start+len(df)
    for col in scale_cols:
        L=get_column_letter(list(df.columns).index(col)+1)
        ws.conditional_formatting.add(f"{L}{start+1}:{L}{last}",
            ColorScaleRule(start_type="min",start_color="F8696B",
                           mid_type="percentile",mid_value=55,mid_color="FFEB84",
                           end_type="max",end_color="63BE7B"))
    if width_map: widths(ws,width_map)
    ws.freeze_panes=ws[f"A{start+1}"]
    if img:
        im=XLImage(img); im.width=470; im.height=360; ws.add_image(im,img_anchor or f"{get_column_letter(len(df.columns)+2)}2")
    return ws

# ---- Sheet 2: CHAMPION FORECAST (full 48) -------------------------------- #
cf=full[["rank","team","group","confederation","elo","P_champion","champ_SE","P_market",
         "P_consensus","value_vs_market","P_reach_Final","P_reach_SF","P_reach_QF",
         "P_reach_R16","P_reach_R32"]].copy()
cf.columns=["Rank","Team","Grp","Confed","Elo","Model","±SE","Market","Consensus","Value(M−Mkt)",
            "Final","Semi","Quarter","R16","R32"]
add_df_sheet("Champion Forecast",cf,
    pct_cols=["Model","±SE","Market","Consensus","Value(M−Mkt)","Final","Semi","Quarter","R16","R32"],
    int_cols=["Rank","Elo"],
    scale_cols=["Model","Consensus","Final","Semi","Quarter","R16","R32"],
    width_map={"A":5,"B":18,"C":5,"D":12,"E":7,"F":9,"G":7,"H":9,"I":11,"J":12,
               "K":8,"L":8,"M":9,"N":8,"O":8},
    title="CHAMPION FORECAST — all 48 teams (sorted by consensus title probability)")

# ---- Sheet 3: GROUP STAGE ------------------------------------------------- #
gs=grp[["group","team","elo","E_pts_mc","P_win_group","P_runner_up","P_third","P_top2","P_reach_R32"]].copy()
gs.columns=["Grp","Team","Elo","Exp.Pts","Win Group","Runner-up","3rd","Top-2","Reach R32"]
add_df_sheet("Group Stage",gs,
    pct_cols=["Win Group","Runner-up","3rd","Top-2","Reach R32"],int_cols=["Elo"],
    scale_cols=["Win Group","Top-2","Reach R32"],
    width_map={"A":5,"B":18,"C":7,"D":9,"E":11,"F":11,"G":8,"H":9,"I":10},
    title="GROUP STAGE — qualification probabilities (Monte-Carlo, tiebreakers applied)")

# ---- Sheet 4: APPENDIX A (Poisson-Binomial) ------------------------------ #
aa=appA.copy()
aa.columns=["Grp","Team","Elo","Win m1","Win m2","Win m3","Draw avg","Loss avg","E[Pts]",
            "P(0W)","P(1W)","P(2W)","P(3W)"]+[f"P(pts={k})" for k in range(10)]
add_df_sheet("Appendix A — PoissonBinomial",aa,
    pct_cols=["Win m1","Win m2","Win m3","Draw avg","Loss avg","P(0W)","P(1W)","P(2W)","P(3W)"]+
             [f"P(pts={k})" for k in range(10)],
    int_cols=["Elo"], scale_cols=["P(3W)","P(pts=9)"],
    width_map={"A":4,"B":17,"C":6},
    title="TECHNICAL APPENDIX A — exact per-team Poisson-binomial win distribution & points distribution")

# ---- Sheet 5: MATCH MATRIX ------------------------------------------------ #
mmx=mm.copy()
mmx.columns=["Grp","Team A","Team B","P(A win)","P(draw)","P(B win)","xG A","xG B"]
add_df_sheet("Match Matrix",mmx,
    pct_cols=["P(A win)","P(draw)","P(B win)"],
    scale_cols=["P(A win)","P(B win)"],
    width_map={"A":4,"B":17,"C":17,"D":9,"E":9,"F":9,"G":7,"H":7},
    title="GROUP-STAGE MATCH MATRIX — 72 fixtures (W/D/L probabilities and expected goals)")

# ---- Sheet 6: MODEL vs MARKET -------------------------------------------- #
mv=full[["rank","team","P_champion","P_market","P_consensus","value_vs_market"]].copy()
mv=mv.sort_values("value_vs_market",ascending=False)
mv.columns=["Rank","Team","Model","Market","Consensus","Value (Model−Market)"]
add_df_sheet("Model vs Market",mv,
    pct_cols=["Model","Market","Consensus","Value (Model−Market)"],int_cols=["Rank"],
    scale_cols=["Value (Model−Market)"],
    width_map={"A":5,"B":18,"C":9,"D":9,"E":11,"F":18},
    title="MODEL vs MARKET — value board (top = model backs; bottom = model fades)",
    img=f"{FIG}/fig2_value_vs_market.png",img_anchor="H2")

# ---- Sheet 7: SENSITIVITY ------------------------------------------------- #
add_df_sheet("Sensitivity",sens.rename(columns={"form_sd":"sigma_form (Elo)"}),
    pct_cols=[], width_map={"A":16,"B":8,"C":10,"D":8,"E":9,"F":9,"G":8},
    title="SENSITIVITY — title odds (%) vs the form-uncertainty parameter",
    img=f"{FIG}/fig6_sensitivity.png",img_anchor="I2")

# ---- Sheet 8: METHODOLOGY & SOURCES -------------------------------------- #
ws=wb.create_sheet("Methodology & Sources"); ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=120
ws.merge_cells("B2:B2"); ws["B2"].value="METHODOLOGY, ASSUMPTIONS & SOURCES"
ws["B2"].font=Font(size=14,bold=True,color="0A1F44")
method_text=[
 "DEVELOPMENT CHAIN",
 "1. Strength rating R_i = World Football Elo (eloratings.net, current mid-June 2026 snapshot).",
 "2. Elo→expected score:  E_A = 1 / (1 + 10^(−(R_A−R_B)/400)).",
 "3. Goal model: independent Poisson; total τ(d)=2.70+0.0012·|d|, supremacy σ(d) solved so the",
 "   Skellam win/draw/loss reproduces E_A exactly (calibration max error 0.4%).",
 "4. Group stage solved EXACTLY: each team's win count is Poisson-binomial(p1,p2,p3); the full",
 "   points distribution from enumerating all 3^6=729 result combinations (verified vs DFT & MC).",
 "5. Knockout: 200,000 Monte-Carlo replications over FIFA's fixed Round-of-32 bracket, including the",
 "   best-third-place 8-of-12 combination assignment (bipartite SDR; all 495 patterns valid).",
 "6. Extra time = goal rate × 30/90; shootout ≈ coin flip with a tiny Elo edge (±0.12 max).",
 "7. Uncertainty: per-tournament team 'form' random effect σ=70 Elo (rating error + injuries + draw luck).",
 "   Host advantage +65 Elo for USA / Mexico / Canada. Seed=2026 (fully reproducible).",
 "",
 "KEY ASSUMPTIONS / LIMITATIONS",
 "• Elo encodes results only — not roster ageing, injuries, or tactical match-ups (hence form σ & market blend).",
 "• Independent Poisson (no explicit Dixon-Coles correlation) keeps the exact Elo calibration intact.",
 "• Third-place routing follows FIFA eligibility constraints; exact 495-row table has negligible effect on title odds.",
 "• Monte-Carlo standard error on a 20% probability at N=200k is ≈0.09 percentage points.",
 "",
 "PRIMARY DATA SOURCES",
 "• World Football Elo Ratings — eloratings.net (World.tsv, retrieved 17 Jun 2026).",
 "• FIFA Men's World Ranking, 11 Jun 2026 (cross-reference).",
 "• 2026 FIFA World Cup final draw — Wikipedia & simbye.com (groups cross-validated).",
 "• Knockout bracket — FIFA fixed match schedule (Wikipedia 2026 WC knockout stage).",
 "• Market odds — FOX Sports / consensus sportsbooks, 17 Jun 2026 (de-vigged).",
 "• Benchmarks — Nate Silver 'PELE' model; Kimi 300-agent qualitative forecast (this repo).",
 "",
 "SELECTED LITERATURE (modelling grounding)",
 "• Dixon & Coles (1997) Applied Statistics 46(2) — bivariate Poisson football scores.",
 "• Hvattum & Arntzen (2010) Int. J. Forecasting — Elo ratings for football forecasting.",
 "• Maher (1982) Statistica Neerlandica — Poisson models for football.",
 "• Le Cam (1960) — Poisson-binomial / Poisson approximation theory.",
]
r=4
for line in method_text:
    c=ws.cell(row=r,column=2,value=line)
    if line.isupper() and line.strip(): c.font=Font(bold=True,size=11,color="0A1F44")
    else: c.font=Font(size=10)
    c.alignment=Alignment(wrap_text=True,vertical="top"); r+=1

# embed calibration + heatmap images on methodology sheet
im=XLImage(f"{FIG}/fig4_calibration.png"); im.width=430; im.height=300; ws.add_image(im,"D4")
im2=XLImage(f"{FIG}/fig3_round_heatmap.png"); im2.width=360; im2.height=400; ws.add_image(im2,"D22")

# native Excel bar chart on Champion Forecast (top 12)
wsc=wb["Champion Forecast"]
chart=BarChart(); chart.type="bar"; chart.title="Title probability — top 12"
chart.height=9; chart.width=16
data=Reference(wsc,min_col=6,min_row=2,max_row=14)         # Model column
cats=Reference(wsc,min_col=2,min_row=3,max_row=14)
chart.add_data(data,titles_from_data=True); chart.set_categories(cats)
chart.y_axis.numFmt="0%"; chart.legend=None
wsc.add_chart(chart,"Q2")

path=os.path.join(OUT,"WorldCup2026_Forecast.xlsx")
wb.save(path)
print("workbook saved:",path)
print("sheets:",wb.sheetnames)
